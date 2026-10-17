"""后端能力自检：文档解析（Word / Markdown）、切片、ChromaDB 向量检索、ORM 模型。

不依赖 MySQL 与智谱 API Key，可独立运行：
    .venv\\Scripts\\python.exe verify_capabilities.py

会被验证的能力：
  1. Word（.docx）解析能取到段落与表格
  2. Markdown（.md）解析能按标题切节
  3. 切片参数与 rag 模块一致（500/50），元数据保留页码
  4. ChromaDB 能写入并检索到内容，且 owner 隔离生效
  5. ORM 模型能被 SQLAlchemy 正确编译（含 MySQL 方言）
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

RESULTS: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, note: str = "") -> None:
    RESULTS.append((name, bool(ok), note))


def section(title: str) -> None:
    print(f"\n{'=' * 70}\n{title}\n{'=' * 70}")


# --------------------------------------------------------------------------
section("1. 依赖导入")
# --------------------------------------------------------------------------
try:
    import chromadb  # noqa: F401

    check("chromadb 可导入", True, f"版本 {chromadb.__version__}")
except Exception as exc:
    check("chromadb 可导入", False, str(exc))

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter  # noqa: F401

    check("langchain-text-splitters 可导入", True)
except Exception as exc:
    check("langchain-text-splitters 可导入", False, str(exc))

try:
    import aiomysql  # noqa: F401

    check("aiomysql（MySQL 驱动）可导入", True)
except Exception as exc:
    check("aiomysql（MySQL 驱动）可导入", False, str(exc))

try:
    from app.services import chunking, parser, vectorstore

    check("后端服务模块可导入", True)
except Exception as exc:
    check("后端服务模块可导入", False, f"{type(exc).__name__}: {exc}")
    for name, ok, note in RESULTS:
        print(f"{'PASS' if ok else 'FAIL'}  {name}  {note}")
    sys.exit(1)


# --------------------------------------------------------------------------
section("2. Word（.docx）解析")
# --------------------------------------------------------------------------
tmp = Path(tempfile.mkdtemp(prefix="czb-verify-"))

try:
    import docx

    document = docx.Document()
    document.add_heading("创智杯测试文档", level=1)
    document.add_paragraph("这是一段正文，用于验证 Word 解析能正确提取段落。")
    document.add_heading("表格测试", level=2)
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "步骤"
    table.cell(0, 1).text = "责任人"
    table.cell(0, 2).text = "时限"
    table.cell(1, 0).text = "提交票据"
    table.cell(1, 1).text = "申请人"
    table.cell(1, 2).text = "15 日内"
    docx_path = tmp / "word-test.docx"
    document.save(str(docx_path))

    documents = parser.parse_document(docx_path)
    text = "\n".join(d.page_content for d in documents)

    check("Word 文件解析成功", bool(documents), f"{len(documents)} 个 Document")
    check("Word 正文被提取", "验证 Word 解析能正确提取段落" in text, "")
    check(
        "Word 表格被提取",
        "提交票据" in text and "申请人" in text,
        "表格内容未丢失" if "提交票据" in text else "表格内容丢失",
    )
except Exception as exc:
    check("Word 文件解析", False, f"{type(exc).__name__}: {exc}")


# --------------------------------------------------------------------------
section("3. Markdown（.md）解析")
# --------------------------------------------------------------------------
try:
    md_path = tmp / "markdown-test.md"
    md_path.write_text(
        "# 系统架构设计\n\n"
        "## 整体架构\n"
        "前端 Vue 3 加 Vite，后端 Python 加 FastAPI。\n\n"
        "## 文本分块\n"
        "块大小默认 500 字，重叠 50 字。\n\n"
        "## 向量化\n"
        "使用智谱 embedding-3，存入 ChromaDB。\n",
        encoding="utf-8",
    )

    documents = parser.parse_document(md_path)
    text = "\n".join(d.page_content for d in documents)

    check("Markdown 解析成功", bool(documents), f"{len(documents)} 个 Document")
    check("Markdown 按标题切节", len(documents) >= 3, f"切出 {len(documents)} 节")
    check("Markdown 中文内容完整", "智谱 embedding-3" in text, "")
except Exception as exc:
    check("Markdown 解析", False, f"{type(exc).__name__}: {exc}")


# --------------------------------------------------------------------------
section("4. 切片（参数与 rag 模块一致）")
# --------------------------------------------------------------------------
try:
    from app.core.config import settings

    check("切片长度为 500（对齐 rag/build_db.py）", settings.CHUNK_SIZE == 500,
          f"当前 {settings.CHUNK_SIZE}")
    check("切片重叠为 50", settings.CHUNK_OVERLAP == 50, f"当前 {settings.CHUNK_OVERLAP}")

    long_md = tmp / "long.md"
    long_md.write_text(
        "# 长文档测试\n\n" + "这是一段用于验证切片的中文文本。" * 200, encoding="utf-8"
    )
    documents = parser.parse_document(long_md)
    chunks = chunking.split_documents(documents)

    check("长文档被切成多块", len(chunks) > 1, f"{len(chunks)} 块")
    check(
        "切片长度不超过上限（含少量分隔符余量）",
        all(len(c.page_content) <= settings.CHUNK_SIZE + 60 for c in chunks),
        f"最长 {max((len(c.page_content) for c in chunks), default=0)}",
    )
    check(
        "切片保留来源元数据",
        all((c.metadata or {}).get("source") for c in chunks),
        "",
    )
except Exception as exc:
    check("切片流程", False, f"{type(exc).__name__}: {exc}")


# --------------------------------------------------------------------------
section("5. ChromaDB 写入与检索")
# --------------------------------------------------------------------------
try:
    from app.services import vectorstore

    health = vectorstore.health()
    check("向量库可用", health.get("available") is True, str(health.get("error", "")))
    check("集合名带向量空间后缀", "zhipu2048" in health.get("collection", "")
          or "default" in health.get("collection", ""), health.get("collection", ""))
    print(f"    向量模式：{health.get('embedding', {}).get('mode')}")
    print(f"    集合名称：{health.get('collection')}")

    owner_a = "user_verify_a"
    owner_b = "user_verify_b"
    kb_id = "kb_verify"

    written = vectorstore.add_chunks(
        owner_id=owner_a,
        knowledge_base_id=kb_id,
        knowledge_base_name="自检知识库",
        document_id="doc_verify",
        document_name="自检文档.md",
        chunks=[
            (0, "蓝色鲸鱼协议编号是 CZ-2025，用于系统联调测试。", 1),
            (1, "本系统的向量库使用 ChromaDB 存储，向量模型为智谱 embedding-3。", 2),
            (2, "报销流程共四步：提交票据、主管审批、财务复核、打款。", 3),
        ],
    )
    check("向量写入成功", written == 3, f"写入 {written} 条")

    hits = vectorstore.search(
        None, "蓝色鲸鱼协议的编号是多少", owner_id=owner_a, knowledge_base_ids=[kb_id], top_k=3
    )
    check("检索能召回写入内容", len(hits) > 0, f"命中 {len(hits)} 条")
    if hits:
        top = hits[0]
        check(
            "Top-1 命中正确片段",
            "蓝色鲸鱼" in top.content,
            f"Top-1: {top.content[:30]}（相似度 {top.score}）",
        )
        check("检索结果带回文档名", bool(top.document_name), top.document_name)
        check("检索结果带回页码", top.page == 1, f"page={top.page}")

    # 租户隔离：另一个用户不应检索到上面的内容
    other_hits = vectorstore.search(
        None, "蓝色鲸鱼协议的编号是多少", owner_id=owner_b, top_k=3
    )
    check("跨用户隔离生效（他人检索不到）", len(other_hits) == 0, f"他人命中 {len(other_hits)} 条")

    # 知识库范围过滤
    scoped = vectorstore.search(
        None, "蓝色鲸鱼协议", owner_id=owner_a, knowledge_base_ids=["kb_other"], top_k=3
    )
    check("知识库范围过滤生效", len(scoped) == 0, f"范围外命中 {len(scoped)} 条")

    counts = vectorstore.count_by_knowledge_base([kb_id])
    check("按知识库统计切片数正确", counts.get(kb_id) == 3, f"统计 {counts.get(kb_id)}")

    removed = vectorstore.delete_document("doc_verify")
    check("删除文档向量成功", removed == 3, f"删除 {removed} 条")
    remaining = vectorstore.count_chunks(knowledge_base_id=kb_id)
    check("删除后统计归零", remaining == 0, f"剩余 {remaining}")
except Exception as exc:
    import traceback

    check("ChromaDB 写入与检索", False, f"{type(exc).__name__}: {exc}")
    traceback.print_exc()


# --------------------------------------------------------------------------
section("6. ORM 模型与 MySQL 方言")
# --------------------------------------------------------------------------
try:
    from sqlalchemy.schema import CreateTable
    from sqlalchemy.dialects import mysql

    from app import models

    check("模型加载成功", True, f"{len(models.Base.metadata.tables)} 张表")
    check(
        "已移除 Chunk 表（向量交给 ChromaDB）",
        "chunks" not in models.Base.metadata.tables,
        "表清单：" + ", ".join(sorted(models.Base.metadata.tables)),
    )

    ddl = str(CreateTable(models.Document.__table__).compile(dialect=mysql.dialect()))
    check("MySQL 方言可编译 DDL", "CREATE TABLE documents" in ddl, "")
    check(
        "文本列使用 utf8mb4 兼容类型",
        "VARCHAR" in ddl and "TEXT" in ddl,
        "",
    )
    check(
        "索引列长度安全（utf8mb4 下不超 255）",
        all(
            (col.type.length or 0) <= 255
            for table in models.Base.metadata.tables.values()
            for col in table.columns
            if col.index and hasattr(col.type, "length")
        ),
        "",
    )

    from app.core.config import settings as cfg

    check("DSN 指向 MySQL", cfg.mysql_dsn.startswith("mysql+aiomysql"), cfg.mysql_dsn.split("@")[-1])
except Exception as exc:
    import traceback

    check("ORM 与 MySQL 方言", False, f"{type(exc).__name__}: {exc}")
    traceback.print_exc()


# --------------------------------------------------------------------------
print(f"\n{'=' * 70}")
print("自检结果")
print("=" * 70)
for name, ok, note in RESULTS:
    suffix = f"  ← {note}" if note else ""
    print(f"{'PASS' if ok else 'FAIL'}  {name}{suffix}")
failed = [r for r in RESULTS if not r[1]]
print("-" * 70)
print(f"{len(RESULTS) - len(failed)}/{len(RESULTS)} 通过")
sys.exit(1 if failed else 0)
